from openpyxl import Workbook


class ExcelExporter:
    def export(self, report, filename="plan.xlsx"):
        wb = Workbook()

        # Ark 1: Plan
        ws = wb.active
        ws.title = "Plan"

        ws.append([
            "Installation",
            "Aktivitet",
            "Hold",
            "Start",
            "Slut",
        ])

        for aktivitet in report.schedule_result.activities:
            ws.append([
                aktivitet.installation_id,
                aktivitet.type,
                aktivitet.hold,
                aktivitet.start_dato,
                aktivitet.slut_dato,
            ])

        # Ark 2: Restarbejde
        ws = wb.create_sheet("Restarbejde")

        ws.append([
            "Installation",
            "Aktivitet",
            "Antal",
            "Årsag",
        ])

        for rest in report.schedule_result.rest_work:
            ws.append([
                rest.installation_id,
                rest.task_type,
                rest.quantity,
                rest.reason,
            ])

        # Ark 3: Genåbningsforslag
        ws = wb.create_sheet("Genåbningsforslag")

        ws.append([
            "Zone",
            "Start",
            "Slut",
            "Varighed",
        ])

        for proposal in report.reopen_proposals:
            ws.append([
                proposal.zone,
                proposal.suggested_start,
                proposal.suggested_end,
                proposal.estimated_days_total,
            ])

        # Ark 4: Genåbningsplan
        ws = wb.create_sheet("Genåbningsplan")

        ws.append([
            "Hold",
            "Aktivitet",
            "Start",
            "Slut",
            "Antal",
        ])

        for schedule in report.reopen_schedules:
            for activity in schedule.activities:
                ws.append([
                    activity.hold,
                    activity.task_type,
                    activity.start_dato,
                    activity.slut_dato,
                    activity.quantity,
                ])

                # Ark pr. hold
        if hasattr(report, "hold_reports"):
            for hold_report in report.hold_reports:
                ws = wb.create_sheet(f"Hold {hold_report.hold}")

                ws.append([
                    "Dato fra",
                    "Dato til",
                    "Installation",
                    "Aktivitet",
                ])

                for aktivitet in sorted(
                    hold_report.aktiviteter,
                    key=lambda x: x.start_dato
                ):
                    ws.append([
                        aktivitet.start_dato,
                        aktivitet.slut_dato,
                        aktivitet.installation_id,
                        aktivitet.aktivitetstype,
                    ])
                # Ark pr. hold
        for hold_report in report.hold_reports:
            sheet_name = f"Hold {hold_report.hold}"[:31]
            ws = wb.create_sheet(sheet_name)

            ws.append([
                "Dato fra",
                "Dato til",
                "Hold",
                "Installation",
                "Aktivitet",
            ])

            for aktivitet in sorted(
                hold_report.aktiviteter,
                key=lambda x: x.start_dato
            ):
                ws.append([
                    aktivitet.start_dato,
                    aktivitet.slut_dato,
                    aktivitet.hold,
                    aktivitet.installation_id,
                    aktivitet.aktivitetstype,
                ])

        # Ark: Holdbelægning
        if getattr(report, "resource_report", None):
            ws = wb.create_sheet("Holdbelægning")

            ws.append([
                "Hold",
                "År",
                "Måned",
                "Arbejdsdage",
                "Bookede dage",
                "Belægning %",
            ])

            for item in sorted(
                report.resource_report.months,
                key=lambda x: (x.hold, x.year, x.month)
            ):
                ws.append([
                    item.hold,
                    item.year,
                    item.month,
                    item.available_days,
                    item.booked_days,
                    item.utilization_percent,
                ])

        wb.save(filename)

        print(f"\n📄 Excel gemt: {filename}")
